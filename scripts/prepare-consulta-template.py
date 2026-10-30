"""Editable PDI form positioned against the user's photographed sheet.

The photo controls layout and fixed wording. The user requested the sample
PDI officer block and background stamp; the handwritten signature is excluded.
Run with the Codex bundled Python runtime (python-docx).
"""
from copy import deepcopy
from pathlib import Path
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'assets/templates/pdi-consulta-v1.xml'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Pt(612), Pt(936)
sec.top_margin, sec.bottom_margin = Pt(28), Pt(36)
sec.left_margin, sec.right_margin = Pt(68), Pt(32)
style = doc.styles['Normal']
style.font.name, style.font.size = 'Arial', Pt(11)
style.paragraph_format.space_after = Pt(0)
LINE, BODY, RIGHT = 13.25, 143, 567
y = 28
checkbox_id = 960000

def run(p, text, bold=False, underline=False, italic=False, size=11, caps=False):
    r = p.add_run(text)
    r.bold, r.underline, r.italic = bold, underline, italic
    r.font.name, r.font.size, r.font.all_caps = 'Arial', Pt(size), caps
    return r

def checkbox(p, tag, title, checked=False):
    """Native Word checkbox: clickable without macros or form protection.

    Set the symbol font explicitly for every script; inheriting Arial from the
    acta can render these Unicode boxes as unrelated fallback glyphs in Word.
    """
    global checkbox_id
    checkbox_id += 1
    sdt = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr')
    for name, value in [('alias', title), ('tag', tag), ('id', str(checkbox_id))]:
        element = OxmlElement('w:' + name)
        element.set(qn('w:val'), value)
        props.append(element)
    control = OxmlElement('w14:checkbox')
    state = OxmlElement('w14:checked')
    state.set(qn('w14:val'), '1' if checked else '0')
    control.append(state)
    for name, value in [('checkedState', '2612'), ('uncheckedState', '2610')]:
        symbol = OxmlElement('w14:' + name)
        symbol.set(qn('w14:val'), value)
        symbol.set(qn('w14:font'), 'Segoe UI Symbol')
        control.append(symbol)
    props.append(control)
    sdt.append(props)
    content = OxmlElement('w:sdtContent')
    r = run(p, '☒' if checked else '☐')
    fonts = r._r.get_or_add_rPr().rFonts
    for script in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + script), 'Segoe UI Symbol')
    fonts.set(qn('w:hint'), 'default')
    content.append(r._r)
    sdt.append(content)
    p._p.append(sdt)

def background_stamp(p):
    """Floating image behind editable officer text, without changing pagination.

    The application attaches the referenced PNG to the final DOCX package.
    All positions are in page coordinates, converted from points to EMU.
    """
    from PIL import Image
    with Image.open(ROOT / 'assets/templates/pdi-sello.png') as stamp:
        width, height = 124, 124 * stamp.height / stamp.width
    emu = lambda pt: str(round(pt * 12700))
    drawing = etree.fromstring(f'''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page"><wp:posOffset>{emu(370)}</wp:posOffset></wp:positionH>
        <wp:positionV relativeFrom="page"><wp:posOffset>{emu(650)}</wp:posOffset></wp:positionV>
        <wp:extent cx="{emu(width)}" cy="{emu(height)}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
        <wp:docPr id="903" name="Sello PDI de referencia" descr="Sello de fondo del bloque PDI"/>
        <wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
        <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
          <pic:pic><pic:nvPicPr><pic:cNvPr id="903" name="pdi-sello.png"/><pic:cNvPicPr/></pic:nvPicPr>
            <pic:blipFill><a:blip r:embed="rIdPdiReferenceStamp"><a:alphaModFix amt="55000"/></a:blip><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
            <pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu(width)}" cy="{emu(height)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>
          </pic:pic>
        </a:graphicData></a:graphic>
      </wp:anchor>
    </w:drawing>''')
    p.add_run()._r.append(drawing)

def paragraph(top, text='', *, x=BODY, right=RIGHT, bold=False, underline=False,
              size=11, center=False, height=LINE, caps=False):
    global y
    if top < y - .1:
        raise ValueError(f'Overlapping layout: {top} < {y}: {text}')
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.left_indent, fmt.right_indent = Pt(x - 68), Pt(580 - right)
    fmt.space_before, fmt.space_after = Pt(top - y), Pt(0)
    fmt.line_spacing = Pt(height)
    fmt.widow_control, fmt.keep_with_next = False, False
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text: run(p, text, bold=bold, underline=underline, size=size, caps=caps)
    y = top + height
    return p

def table(top, rows, widths, *, border=False, row_height=LINE, value_caps=False):
    global y
    if top > y: paragraph(y, height=top-y)
    t = doc.add_table(rows=0, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = t._tbl.tblPr
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), str(int((BODY-68)*20))); ind.set(qn('w:type'), 'dxa'); pr.append(ind)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        item = OxmlElement('w:'+edge)
        item.set(qn('w:val'), 'single' if border else 'nil')
        item.set(qn('w:sz'), '4'); item.set(qn('w:color'), '000000'); borders.append(item)
    pr.append(borders)
    for col, width in zip(t.columns, widths): col.width = Pt(width)
    for data in rows:
        row = t.add_row()
        row.height, row.height_rule = Pt(row_height), WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for index, (cell, width, text) in enumerate(zip(row.cells, widths, data)):
            cell.width = Pt(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top',0),('bottom',0),('left',80 if border else 0),('right',30)]:
                item = OxmlElement('w:'+edge)
                item.set(qn('w:w'), str(value)); item.set(qn('w:type'), 'dxa'); margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0), Pt(0)
            p.paragraph_format.line_spacing = Pt(LINE)
            p.paragraph_format.widow_control = False
            run(p, text, caps=value_caps and index == 1)
    y = top + row_height * len(rows)
    return t

paragraph(28, 'POLICÍA DE INVESTIGACIONES DE CHILE', x=68, right=298, center=True, bold=True)
paragraph(41.25, 'Prefectura de Migraciones y Policía', x=68, right=298, center=True, bold=True, underline=True)
paragraph(54.5, 'Internacional Metropolitana', x=68, right=298, center=True, bold=True, underline=True)
paragraph(82, 'REPORTE DE CONSULTA DE EXTRANJERO N° __________ /', center=True, bold=True)
paragraph(109, 'TIPO DE CONSULTA:', bold=True, underline=True)
p = paragraph(122.25)
p.paragraph_format.tab_stops.add_tab_stop(Pt(254))
run(p, '12° BIS, LEY N° 20.930\t')
checkbox(p, 'pdi_consulta_12bis', 'Consulta por artículo 12 bis', checked=True)
p = paragraph(135.5)
p.paragraph_format.tab_stops.add_tab_stop(Pt(254))
run(p, 'POR DETENCIÓN\t')
checkbox(p, 'pdi_consulta_detencion', 'Consulta por detención')
paragraph(162, 'ANTECEDENTES PERSONALES', bold=True, underline=True)
personal_table = table(188.5, [
    ['PORTA DOC. DE IDENTIDAD', ''],
    ['{{SURNAME_LABEL}}', ': {{PATERNAL_SURNAME}}'],
    ['APELLIDO MATERNO', ': {{MATERNAL_SURNAME}}'],
    ['NOMBRES', ': {{FIRST_NAMES}}'],
    ['NACIONALIDAD', ': {{NATIONALITY}}'],
    ['TIPO DOCUMENTO', ': {{DOCUMENT_TYPE}}'],
    ['N° DOCUMENTO', ': {{DOCUMENT_NUMBER}}'],
    ['FECHA DE NAC.', ': {{BIRTH_DATE}}'],
    ['SEXO', ': {{SEX}}'],
], [170,254], value_caps=True)
p = personal_table.cell(0, 1).paragraphs[0]
run(p, ': SI ')
checkbox(p, 'pdi_porta_documento_si', 'Porta documento de identidad: sí')
run(p, '   NO ')
checkbox(p, 'pdi_porta_documento_no', 'Porta documento de identidad: no')

p = paragraph(323, right=580)
p.paragraph_format.tab_stops.add_tab_stop(Pt(179))
run(p, 'IMPORTANTE :\t', size=10.5)
run(p, 'Los datos señalados ', italic=True, size=10.5)
run(p, 'NO CERTIFICAN', bold=True, italic=True, size=10.5)
run(p, ' la identidad del extranjero.', italic=True, size=10.5)
p = paragraph(336.25, x=220, right=580)
run(p, 'Adicionalmente, si el extranjero ', italic=True, size=10.5)
run(p, 'NO PORTA', bold=True, italic=True, size=10.5)
run(p, ' su documento de identidad, los', italic=True, size=10.5)
p = paragraph(349.5, x=220, right=580)
run(p, 'datos consignados son ', italic=True, size=10.5)
run(p, 'DECLARATIVOS.', bold=True, italic=True, underline=True, size=10.5)
paragraph(376, 'ANTECEDENTES MIGRATORIOS', bold=True, underline=True)
table(400, [['MOVIMIENTO','FECHA','AVANZADA'],['','','']], [98,96,230], border=True, row_height=19)
migration_table = table(451.25, [
    ['SITUACIÓN MIGRATORIA', ''],
    ['MOTIVO', ':'],
    ['REGISTRA DENUNCIA', ':'],
    ['CORRESPONDE DENUNCIA', ':'],
    ['EXPULSIÓN VIGENTE', ':'],
    ['UNIDAD QUE CONSULTA', ': SUBCOMISARIA DE ASUNTOS MIGRATORIOS'],
    ['TIPO DOC.', ': ACTA CONTROL IDENTIDAD'],
    ['NÚMERO DOC.', ': -'],
    ['FECHA', ': {{CONTROL_DATE}}'],
    ['OBSERVACIONES', ': {{NOTES}}'],
], [170,254], value_caps=True)
p = migration_table.cell(0, 1).paragraphs[0]
run(p, ': REGULAR ')
checkbox(p, 'pdi_situacion_regular', 'Situación migratoria: regular')
run(p, '        IRREGULAR ')
checkbox(p, 'pdi_situacion_irregular', 'Situación migratoria: irregular')

paragraph(599, 'Se solicita que, en caso de producirse una detención, se remita copia del Parte Policial al', right=580, size=10.5)
p = paragraph(612.25, right=580)
p.paragraph_format.tab_stops.add_tab_stop(Pt(322))
run(p, 'correo\t, con el propósito de poner estos', size=10.5)
paragraph(625.5, 'antecedentes en conocimiento de la Autoridad Administrativa.', right=580, size=10.5)

# User-requested sample text remains editable; stamp is a separate background.
p = paragraph(682, 'NOMBRE FUNCIONARIO PDI', x=390, right=580, center=True)
background_stamp(p)
paragraph(695.25, 'INSPECTOR', x=390, right=580, center=True)
paragraph(708.5, 'ENCARGADO DE GUARDIA', x=390, right=580, center=True)
paragraph(754.75, 'ILS')
paragraph(768, 'Distribución:', underline=True)
p = paragraph(781.25, '- SUB. ASUNTOS MIGRATORIOS')
p.paragraph_format.tab_stops.add_tab_stop(Pt(258))
run(p, '\t(1)')
p = paragraph(794.5)
p.paragraph_format.tab_stops.add_tab_stop(Pt(258))
run(p, '- Archivo\t(1) /', underline=True)

# Do not inherit typography from the appended Carabineros template.
for p in doc.element.body.iter(qn('w:p')):
    ppr = p.find(qn('w:pPr'))
    if ppr is None: ppr = OxmlElement('w:pPr'); p.insert(0, ppr)
    for rpr in p.iter(qn('w:rPr')):
        color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); rpr.append(color)
section = deepcopy(sec._sectPr)
stype = OxmlElement('w:type'); stype.set(qn('w:val'), 'nextPage'); section.insert(0, stype)
last = doc.element.body.findall(qn('w:p'))[-1]
last.find(qn('w:pPr')).append(section)
fragment = ''.join(etree.tostring(el, encoding='unicode') for el in doc.element.body if el.tag != qn('w:sectPr'))
OUT.write_text(fragment, encoding='utf-8')
print(OUT)
